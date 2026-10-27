"""Engine tests: real workbooks generated and read back with openpyxl."""

from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pytest

from excel_report_maker.config import load_data, load_template
from excel_report_maker.engine import generate_checklist, generate_matrix, generate_report
from excel_report_maker.errors import DataError, TemplateError

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples"


def _load_workbook(payload: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(payload))


def _checklist_bytes() -> bytes:
    template = load_template("staff_checklist")
    data = json.loads((SAMPLES / "staff_checklist_data.json").read_text(encoding="utf-8"))
    buffer = io.BytesIO()
    generate_checklist(template, data, buffer)
    return buffer.getvalue()


def test_checklist_headers_values_suffix_and_freeze() -> None:
    wb = _load_workbook(_checklist_bytes())
    ws = wb["Admin & Compliance"]
    assert [cell.value for cell in ws[1]] == [
        "Target Position",
        "Candidate Name",
        "Nationality",
        "CV Signed? (TECH-6)",
        "Exclusivity Signed?",
        "Code of Conduct Signed? (TECH-7)",
        "SEA/SH Signed? (TECH-8)",
        "Degree Scan",
        "Cert Scan (Critical)",
        "Remuneration",
        "E-mail",
        "Status / Compliance Notes",
    ]
    # Suffix applied once to the first column.
    assert ws["A2"].value == "Position 1: Team Leader:   "
    assert ws["B2"].value == "Alex Johnson"
    assert ws["K2"].value == "alex.johnson@example.com"
    assert ws.freeze_panes == "C2"
    assert ws.max_row == 4  # header + 3 rows


def test_checklist_table_and_conditional_formats_present() -> None:
    wb = _load_workbook(_checklist_bytes())
    ws = wb["Admin & Compliance"]
    assert len(ws.tables) >= 1
    rules = list(ws.conditional_formatting._cf_rules.values())
    assert rules, "expected conditional formatting rules in the workbook"
    assert sum(len(ruleset) for ruleset in rules) >= 3


def test_master_reference_config_features_all_apply() -> None:
    """The v1 showcase config: CF, widths, styles, and page setup now work."""
    template = load_template("master_reference_config")
    buffer = io.BytesIO()
    generate_checklist(template, template["data"], buffer)
    wb = _load_workbook(buffer.getvalue())
    ws = wb["Feature Showcase"]
    # All three conditional-format groups survive (v1 produced zero).
    rules = list(ws.conditional_formatting._cf_rules.values())
    assert sum(len(ruleset) for ruleset in rules) == 3
    # Explicit widths from column_formats are honored (v1 ignored them).
    # openpyxl reports XlsxWriter widths with Excel's ~0.71 padding added.
    assert 10 < ws.column_dimensions["A"].width < 11
    assert 20 < ws.column_dimensions["B"].width < 21
    assert 25 < ws.column_dimensions["E"].width < 26
    # Page setup is honored (v1 ignored it).
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 9
    assert ws["A2"].value == "REF-001"


def test_matrix_sections_score_column_and_width_cap() -> None:
    template = load_template("evaluation_matrix")
    data = json.loads((SAMPLES / "evaluation_matrix_data.json").read_text(encoding="utf-8"))
    buffer = io.BytesIO()
    generate_matrix(template, data, buffer)
    wb = _load_workbook(buffer.getvalue())
    ws = wb["Technical Scoring"]
    assert ws["A2"].value == "Component A: Digital Infrastructure Advisors"
    assert ws["A3"].value == "John Archer"
    assert ws["F3"].value == 96
    assert ws["F3"].font.bold is True  # score_column = 5
    assert ws["A3"].font.bold is not True
    assert ws.column_dimensions["G"].width is not None
    assert ws.column_dimensions["G"].width <= 100  # width_cap


def test_matrix_without_score_column_highlights_nothing() -> None:
    template = load_template("evaluation_scorecard")
    del template["score_column"]
    data = json.loads((SAMPLES / "evaluation_matrix_data.json").read_text(encoding="utf-8"))
    buffer = io.BytesIO()
    generate_matrix(template, data, buffer)
    ws = _load_workbook(buffer.getvalue())["Evaluation Scorecard"]
    assert ws["F3"].font.bold is not True


def test_empty_checklist_produces_header_only_workbook() -> None:
    template = load_template("compliance_checklist")
    buffer = io.BytesIO()
    generate_checklist(template, [], buffer)
    ws = _load_workbook(buffer.getvalue())["Compliance Checklist"]
    assert ws.max_row == 1
    assert ws["A1"].value == "Position ID"


def test_empty_matrix_produces_header_only_workbook() -> None:
    template = load_template("evaluation_matrix")
    buffer = io.BytesIO()
    generate_matrix(template, [], buffer)
    ws = _load_workbook(buffer.getvalue())["Technical Scoring"]
    assert ws.max_row == 1


def test_ragged_checklist_row_names_the_row() -> None:
    template = load_template("compliance_checklist")
    with pytest.raises(DataError, match="row 1 has 1 value"):
        generate_checklist(
            template,
            [
                [
                    "POS-1",
                    "Alex Johnson",
                    "Canada",
                    "Complete",
                    "Complete",
                    "Complete",
                    "Complete",
                    "$1",
                    "a@example.com",
                    "OK",
                ],
                ["POS-2"],
            ],
            io.BytesIO(),
        )


def test_ragged_matrix_row_names_section_and_row() -> None:
    template = load_template("evaluation_matrix")
    data = [{"header": "Section", "rows": [["Only Name"]]}]
    with pytest.raises(DataError, match="section 0 row 0"):
        generate_matrix(template, data, io.BytesIO())


def test_none_payloads_are_data_errors() -> None:
    template = load_template("staff_checklist")
    with pytest.raises(DataError):
        generate_checklist(template, None, io.BytesIO())
    matrix = load_template("evaluation_matrix")
    with pytest.raises(DataError):
        generate_matrix(matrix, None, io.BytesIO())


def test_unknown_report_type_and_bad_sheet_name_rejected() -> None:
    bad_type = {"report_type": "pivot", "sheet_name": "S", "columns": ["A"]}
    with pytest.raises(TemplateError, match="report_type"):
        generate_report(bad_type, [], io.BytesIO())
    bad_sheet = {"report_type": "checklist", "sheet_name": "X" * 32, "columns": ["A"]}
    with pytest.raises(TemplateError, match="31-character"):
        generate_report(bad_sheet, [], io.BytesIO())


def test_suffix_not_doubled_and_numbers_untouched() -> None:
    template = {
        "report_type": "checklist",
        "sheet_name": "S",
        "columns": ["Label", "Value"],
        "column_formats": [{"index": 0, "suffix": ": "}],
    }
    buffer = io.BytesIO()
    generate_checklist(template, [["Already: ", 42], ["Plain", 7]], buffer)
    ws = _load_workbook(buffer.getvalue())["S"]
    assert ws["A2"].value == "Already: "
    assert ws["A3"].value == "Plain: "
    assert ws["B2"].value == 42


def test_comment_keys_in_rules_are_stripped_and_unknown_format_fails() -> None:
    base = {
        "report_type": "checklist",
        "sheet_name": "S",
        "columns": ["Status"],
        "conditional_formats": [
            {
                "index": 0,
                "rules": [
                    {
                        "_comment": "documentation only",
                        "type": "text",
                        "criteria": "containing",
                        "value": "Done",
                        "format": "green",
                    }
                ],
            }
        ],
    }
    buffer = io.BytesIO()
    generate_checklist(base, [["Done"]], buffer)
    ws = _load_workbook(buffer.getvalue())["S"]
    assert list(ws.conditional_formatting._cf_rules.values())

    broken = json.loads(json.dumps(base))
    broken["conditional_formats"][0]["rules"][0]["format"] = "no_such_format"
    with pytest.raises(TemplateError, match="unknown format name"):
        generate_checklist(broken, [["Done"]], io.BytesIO())


def test_generate_report_dispatch_and_explicit_path(tmp_path: Path) -> None:
    template = load_template("compliance_checklist")
    data = load_data(SAMPLES / "compliance_checklist_data.json", template)
    out = tmp_path / "out.xlsx"
    result = generate_report(template, data, str(out))
    assert result == str(out)
    assert out.exists()
    ws = openpyxl.load_workbook(out)["Compliance Checklist"]
    assert ws["B2"].value == "Alex Johnson"


def test_matrix_section_without_header_rejected() -> None:
    template = load_template("evaluation_matrix")
    with pytest.raises(DataError, match="header"):
        generate_matrix(template, [{"rows": []}], io.BytesIO())


def test_bad_file_name_rejected() -> None:
    template = load_template("staff_checklist")
    template["file_name"] = "../escape.xlsx"
    with pytest.raises(TemplateError, match="file_name"):
        generate_checklist(template, [], io.BytesIO())
