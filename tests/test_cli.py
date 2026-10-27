"""CLI tests, including the real ``python -m`` entry point."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from excel_report_maker.cli import main

ROOT = Path(__file__).resolve().parents[1]


def test_list_templates(capsys: pytest.CaptureFixture[str]) -> None:
    assert main(["list-templates"]) == 0
    assert "staff_checklist" in capsys.readouterr().out


def test_version_flag_exits_zero() -> None:
    with pytest.raises(SystemExit) as excinfo:
        main(["--version"])
    assert excinfo.value.code == 0


def test_validate_ok_and_invalid_data(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    data = ROOT / "data" / "samples" / "staff_checklist_data.json"
    assert main(["validate", "--template", "staff_checklist", "--data", str(data)]) == 0
    assert "Valid" in capsys.readouterr().out

    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps([["short"]]), encoding="utf-8")
    assert main(["validate", "--template", "staff_checklist", "--data", str(bad)]) == 1


def test_validate_unknown_template_returns_one() -> None:
    assert main(["validate", "--template", "nope"]) == 1


def test_generate_in_process(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    out = tmp_path / "report.xlsx"
    data = ROOT / "data" / "samples" / "evaluation_matrix_data.json"
    assert main(["generate", "--template", "evaluation_matrix", "--data", str(data), "--output", str(out)]) == 0
    assert "Generated" in capsys.readouterr().out
    ws = openpyxl.load_workbook(out)["Technical Scoring"]
    assert ws["A3"].value == "John Archer"


def test_generate_via_module_entry_point(tmp_path: Path) -> None:
    out = tmp_path / "report.xlsx"
    data = ROOT / "data" / "samples" / "staff_checklist_data.json"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "excel_report_maker",
            "generate",
            "--template",
            "staff_checklist",
            "--data",
            str(data),
            "--output",
            str(out),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert out.exists()
    ws = openpyxl.load_workbook(out)["Admin & Compliance"]
    assert ws["B2"].value == "Alex Johnson"
